"""
Universal Excel Parser and Schema Normalizer.
Handles diverse sheets, differing column headers, section headers, and multi-turn queries.
"""
import openpyxl
import re
from typing import List, Dict, Any, Optional

def normalize_header(text: Any) -> str:
    if text is None:
        return ""
    # strip extra whitespace, lowercase, remove punctuation
    cleaned = re.sub(r'[^a-zA-Z0-9\s_]', '', str(text).strip().lower())
    return " ".join(cleaned.split())

class ExcelParser:
    def __init__(self, synonym_map: Dict[str, List[str]]):
        self.synonym_map = {
            canonical: [normalize_header(s) for s in synonyms]
            for canonical, synonyms in synonym_map.items()
        }

    def _match_column(self, cell_val: Any) -> Optional[str]:
        norm = normalize_header(cell_val)
        if not norm:
            return None
        # 1. Exact match first across all canonicals
        for canonical, synonyms in self.synonym_map.items():
            if norm in synonyms:
                return canonical
        # 2. Substring match fallback only if no exact match found
        if len(norm) <= 35:
            for canonical, synonyms in self.synonym_map.items():
                for s in synonyms:
                    if len(s) >= 4 and (s in norm or norm in s):
                        return canonical
        return None

    def parse_sheet(self, filepath: str, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active

        test_cases = []
        current_scenario = "General Scenarios"
        header_map = {} # canonical -> col_idx (1-based)

        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            non_empty = [(c + 1, v) for c, v in enumerate(row_vals) if v is not None and str(v).strip()]

            if not non_empty:
                continue

            # Check if this row is a Section header (e.g., Scenario ID 01 General User)
            row_text = " ".join(str(v).strip() for _, v in non_empty)
            if len(non_empty) <= 3 and any(k in row_text.lower() for k in ["scenario id", "test cases", "user types", "admin panel"]):
                current_scenario = row_text
                continue

            # Check if this row is a Column Header row
            candidate_headers = {}
            for col_idx, val in non_empty:
                canonical = self._match_column(val)
                if canonical:
                    candidate_headers[canonical] = col_idx

            # A valid header row must contain at least 3 distinct recognized header types
            if len(candidate_headers) >= 3 and ("input_data" in candidate_headers or "expected" in candidate_headers):
                header_map = candidate_headers
                continue

            # If headers are established, parse this row as a test case
            if header_map and ("input_data" in header_map or "steps" in header_map):
                input_col = header_map.get("input_data")
                expected_col = header_map.get("expected")
                test_id_col = header_map.get("test_id")
                desc_col = header_map.get("description")
                feature_col = header_map.get("feature")
                type_col = header_map.get("test_type")
                steps_col = header_map.get("steps")
                status_col = header_map.get("test_status")
                result_col = header_map.get("test_result")
                score_col = header_map.get("score")
                remark_col = header_map.get("remark")
                ai_eval_col = header_map.get("ai_evaluation")
                kw_col = header_map.get("expected_keywords")

                raw_input = ws.cell(r, input_col).value if input_col else ""
                raw_expected = ws.cell(r, expected_col).value if expected_col else ""
                raw_test_id = ws.cell(r, test_id_col).value if test_id_col else f"TC-{r}"
                raw_desc = ws.cell(r, desc_col).value if desc_col else ""
                raw_feature = ws.cell(r, feature_col).value if feature_col else ""
                raw_type = ws.cell(r, type_col).value if type_col else "Functional"
                raw_steps = ws.cell(r, steps_col).value if steps_col else ""
                raw_status = ws.cell(r, status_col).value if status_col else ""
                raw_result = ws.cell(r, result_col).value if result_col else ""
                raw_score = ws.cell(r, score_col).value if score_col else ""
                raw_remark = ws.cell(r, remark_col).value if remark_col else ""
                raw_ai_eval = ws.cell(r, ai_eval_col).value if ai_eval_col else "Need AI"
                raw_kw = ws.cell(r, kw_col).value if kw_col else ""

                # If both input_data and expected are empty, skip row
                if not str(raw_input or "").strip() and not str(raw_expected or "").strip():
                    continue

                # In case input_data is empty or N/A, check steps for test query
                actual_query = str(raw_input or "").strip()
                if not actual_query or actual_query.lower() in ["n/a", "none"]:
                    # extract potential query from steps if steps has quotes
                    quotes = re.findall(r'["\']([^"\']{3,})["\']', str(raw_steps or ""))
                    if quotes:
                        actual_query = quotes[0]
                    else:
                        actual_query = str(raw_steps or "").strip()

                turns = self._extract_turns(actual_query)
                # If raw_input only produced 1 turn, but raw_steps contains multi-turn queries, use raw_steps
                if len(turns) <= 1 and raw_steps:
                    steps_turns = self._extract_turns(str(raw_steps))
                    if len(steps_turns) > 1:
                        turns = steps_turns
                        actual_query = str(raw_steps).strip()

                test_cases.append({
                    "row_idx": r,
                    "sheet_name": ws.title,
                    "scenario": current_scenario,
                    "test_id": str(raw_test_id).strip(),
                    "feature": str(raw_feature or "").strip(),
                    "description": str(raw_desc or "").strip(),
                    "test_type": str(raw_type or "").strip(),
                    "status": str(raw_status or "").strip(),
                    "result": str(raw_result or "").strip(),
                    "score": str(raw_score or "").strip(),
                    "input_raw": actual_query,
                    "turns": turns,
                    "expected": str(raw_expected or "").strip(),
                    "remark": str(raw_remark or "").strip(),
                    "ai_evaluation": str(raw_ai_eval or "Need AI").strip(),
                    "expected_keywords": str(raw_kw or "").strip(),
                    "header_map": header_map
                })

        return test_cases

    def _clean_turn_text(self, s: str) -> str:
        s = s.strip()
        if (s.startswith('"') and s.endswith('"')) or (s.startswith("'") and s.endswith("'")):
            s = s[1:-1].strip()
        return s

    def _extract_turns(self, input_text: str) -> List[str]:
        text = str(input_text or "").strip()
        if not text:
            return [""]

        # 1. Pattern: Turn 1: / Turn 1. / Turn 1 - (case-insensitive)
        if re.search(r'\bturn\s*\d+\s*[:\-\.]', text, re.IGNORECASE):
            parts = re.split(r'\bturn\s*\d+\s*[:\-\.]\s*', text, flags=re.IGNORECASE)
            turns = [self._clean_turn_text(p) for p in parts if p.strip()]
            if turns:
                return turns

        # 2. Pattern: Query 1 / Question 1 / Q1 / Step 1 / Prompt 1
        if re.search(r'\b(?:query|question|q|prompt|step)\s*\d+\s*[:\-\.]', text, re.IGNORECASE):
            parts = re.split(r'\b(?:query|question|q|prompt|step)\s*\d+\s*[:\-\.]\s*', text, flags=re.IGNORECASE)
            turns = [self._clean_turn_text(p) for p in parts if p.strip()]
            if turns:
                return turns

        # 3. Pattern: User: ... User: ...
        if len(re.findall(r'\buser\s*[:\-]', text, re.IGNORECASE)) > 1:
            parts = re.split(r'\buser\s*[:\-]\s*', text, flags=re.IGNORECASE)
            turns = [self._clean_turn_text(p) for p in parts if p.strip()]
            if turns:
                return turns

        # 4. Pattern: First: ... Second: ...
        if re.search(r'\bfirst\s*[:\-\.]', text, re.IGNORECASE) and (re.search(r'\bsecond\s*[:\-\.]', text, re.IGNORECASE) or re.search(r'\bthen\s*[:\-\.]', text, re.IGNORECASE)):
            parts = re.split(r'\b(?:first|second|third|fourth|fifth|then|finally)\s*[:\-\.]\s*', text, flags=re.IGNORECASE)
            turns = [self._clean_turn_text(p) for p in parts if p.strip()]
            if turns:
                return turns

        # 5. Pattern: Numbered lines: 1: ... 2: ... or 1. ... 2. ... or 1) ... 2) ...
        if re.search(r'(?:^|\n)\s*1\s*[:\.\)\-]\s*.+(?:^|\n)\s*2\s*[:\.\)\-]', text, re.DOTALL):
            parts = re.split(r'(?:^|\n)\s*\d+\s*[:\.\)\-]\s*', text)
            turns = [self._clean_turn_text(p) for p in parts if p.strip()]
            if len(turns) > 1:
                return turns

        # 6. Fallback: single turn
        clean_text = self._clean_turn_text(text)
        return [clean_text] if clean_text else [""]

